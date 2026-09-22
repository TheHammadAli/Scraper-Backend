"""Excel export: a summary tab, an all-listings tab, and one tab per city."""

from __future__ import annotations

import logging
import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.db import Database

log = logging.getLogger(__name__)

# A short preview sits next to the title so the sheet can be scanned, with the
# complete text kept in its own column further along - the spreadsheet version
# of the click-to-expand row in the web table.
COLUMNS = [
    ("id", "ID", 8),
    ("source", "Source", 12),
    ("source_listing_id", "Listing ID", 16),
    ("city", "City", 14),
    ("area", "Area", 22),
    ("category", "Category", 22),
    ("title", "Title", 45),
    ("description_preview", "Description", 55),
    ("price", "Price", 16),
    ("price_currency", "Currency", 10),
    ("price_raw", "Price (as shown)", 18),
    ("seller_name", "Seller", 20),
    ("phone", "Phone", 16),
    ("ad_date", "Ad Date", 13),
    ("url", "URL", 50),
    ("description", "Full Description", 60),
    ("first_seen_at", "First Seen", 20),
    ("last_seen_at", "Last Seen", 20),
    ("scraped_at", "Scraped At", 20),
]

PREVIEW_CHARS = 90


def _preview(text) -> str:
    """One-line taste of a description, the way the web table shows it."""
    collapsed = " ".join(str(text or "").split())
    if len(collapsed) <= PREVIEW_CHARS:
        return collapsed
    return collapsed[:PREVIEW_CHARS].rstrip() + "..."

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14)
THIN_BORDER = Border(*[Side(style="thin", color="D9D9D9")] * 4)

ILLEGAL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
FORMULA_START = ("=", "+", "-", "@")

# Excel's hard cap is 32767; keep well under it so a huge description cannot
# corrupt the sheet.
MAX_CELL_LENGTH = 30000


def _safe(value):
    """Make a scraped value safe to put in a cell.

    Strips control characters and defuses formula injection - a description
    starting with '=' would otherwise be evaluated when the sheet is opened.
    """
    if value is None or isinstance(value, (int, float)):
        return value
    text = ILLEGAL_CHARS.sub("", str(value))
    if len(text) > MAX_CELL_LENGTH:
        text = text[: MAX_CELL_LENGTH - 3] + "..."
    if text.startswith(FORMULA_START):
        text = "'" + text
    return text


def _sheet_name(name: str) -> str:
    """Excel sheet names: max 31 chars, no []:*?/\\ ."""
    cleaned = re.sub(r"[\[\]:*?/\\]", "-", name).strip() or "Sheet"
    return cleaned[:31]


def _write_header(sheet, row: int = 1) -> None:
    for index, (_, label, width) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=row, column=index, value=label)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", horizontal="left")
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = sheet.cell(row=row + 1, column=1)


def _write_rows(sheet, rows, start_row: int = 2) -> None:
    for offset, record in enumerate(rows):
        row_number = start_row + offset
        for index, (key, _, _) in enumerate(COLUMNS, start=1):
            value = (
                _preview(record["description"])
                if key == "description_preview"
                else record[key]
            )
            cell = sheet.cell(row=row_number, column=index, value=_safe(value))
            cell.border = THIN_BORDER
            # No wrapping: a wrapped description turns every row into a tall
            # block. Each row stays one line, and anyone who wants to read a
            # full description can widen the column or click into the cell.
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            if key == "price" and record[key] is not None:
                cell.number_format = "#,##0"
            elif key == "url" and record[key]:
                cell.hyperlink = record[key]
                cell.font = Font(color="0563C1", underline="single")

    if rows:
        sheet.auto_filter.ref = (
            f"A{start_row - 1}:{get_column_letter(len(COLUMNS))}{start_row + len(rows) - 1}"
        )


def _write_summary(
    sheet,
    db: Database,
    total: int,
    since: str | None = None,
    city: str | None = None,
    seen_since: str | None = None,
) -> None:
    sheet["A1"] = "Listing Collection Summary"
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    sheet["A3"] = f"Total listings: {total:,}"

    notes = []
    if city:
        notes.append(f"city: {city}")
    if since:
        notes.append(f"first seen since {since}")
    if seen_since:
        notes.append(f"collected in the run starting {seen_since}")
    if notes:
        sheet["A4"] = "Filtered to " + ", ".join(notes)

    headers = ["City", "Source", "Listings", "With Price", "With Phone",
               "Avg Price", "Oldest Ad", "Newest Ad"]
    for index, label in enumerate(headers, start=1):
        cell = sheet.cell(row=5, column=index, value=label)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        sheet.column_dimensions[get_column_letter(index)].width = 16

    rows = [
        r
        for r in db.summary_by_city(since=since, seen_since=seen_since)
        if not city or r["city"] == city
    ]
    for offset, record in enumerate(rows):
        row_number = 6 + offset
        values = [record["city"], record["source"], record["listings"],
                  record["with_price"], record["with_phone"], record["avg_price"],
                  record["oldest_ad"], record["newest_ad"]]
        for index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_number, column=index, value=_safe(value))
            cell.border = THIN_BORDER
            if index == 6 and value is not None:
                cell.number_format = "#,##0"

    sheet.freeze_panes = "A6"


def export_to_excel(
    db: Database,
    output_dir: Path,
    filename: str | None = None,
    since: str | None = None,
    city: str | None = None,
    seen_since: str | None = None,
) -> Path:
    """Write the table out, organised city-wise. Returns the file path.

    `since` is an ISO timestamp; when given, only listings first seen at or
    after it are exported - the daily "what's new" view.

    `seen_since` is an ISO timestamp too, but filters on last_seen_at - every
    listing a run touched, new or refreshed. That is what "export what I just
    scraped" means, since an ad collected again is still part of that run.

    `city` narrows the whole workbook to one city, so the export matches the
    filter the caller asked for instead of quietly returning everything.
    """
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    filename = filename or f"listings_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    output_path = output_dir / filename

    all_rows = db.all_listings(city=city, since=since, seen_since=seen_since)
    cities = [city] if city else db.cities_present(since=since, seen_since=seen_since)

    # The listings come first so the file opens on the data. Summary used to
    # be sheet one, which made a freshly downloaded workbook look empty.
    workbook = Workbook()
    combined = workbook.active
    combined.title = "All Listings"
    _write_header(combined)
    _write_rows(combined, all_rows)

    for name in cities:
        city_rows = db.all_listings(city=name, since=since, seen_since=seen_since)
        if not city_rows:
            continue
        sheet = workbook.create_sheet(_sheet_name(name))
        _write_header(sheet)
        _write_rows(sheet, city_rows)

    summary = workbook.create_sheet("Summary")
    _write_summary(
        summary, db, len(all_rows), since=since, city=city, seen_since=seen_since
    )

    workbook.active = 0
    workbook.save(output_path)
    log.info("exported %s listings across %s city sheets -> %s",
             len(all_rows), len(cities), output_path)
    return output_path
