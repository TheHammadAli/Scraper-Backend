"""Offline tests - no network. Run with:  python -m unittest discover tests

Covers the parts that are easy to break silently: price/date/phone parsing,
the dedup key, validation rules, and the Excel export.
"""

from __future__ import annotations

import sys
import tempfile
import unittest
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

from core.db import Database  # noqa: E402
from core.locations import CityOutcome  # noqa: E402
from core.models import Listing, canonical_url  # noqa: E402
from core.normalize import (  # noqa: E402
    extract_phone,
    normalize_phone,
    parse_ad_date,
    parse_price,
)
from export.excel import export_to_excel  # noqa: E402

NOW = datetime(2026, 9, 19, 12, 0, tzinfo=timezone.utc)


def make_listing(**overrides) -> Listing:
    defaults = dict(
        source="olx",
        city="Lahore",
        category="Cars",
        title="Honda Civic 2019 for sale",
        url="https://www.olx.com.pk/item/honda-civic-iid-1234567",
        source_listing_id="1234567",
        description="Excellent condition, single owner, complete documents.",
        price=Decimal("4500000"),
        price_raw="Rs 4,500,000",
    )
    defaults.update(overrides)
    return Listing(**defaults)


class TestCleanText(unittest.TestCase):
    """Zameen stores descriptions with markup, so text has to be normalised."""

    def test_br_tags_become_line_breaks(self):
        from core.normalize import clean_text

        raw = "9 MARLA HOUSE FOR SALE <br /> UNDERGROUND ELECTRICITY <br/>MASJID"
        cleaned = clean_text(raw)
        self.assertNotIn("<", cleaned)
        self.assertNotIn("br", cleaned.lower().replace("marla", ""))
        self.assertIn("UNDERGROUND ELECTRICITY", cleaned)
        self.assertIn("MASJID", cleaned)

    def test_other_tags_are_stripped(self):
        from core.normalize import clean_text

        self.assertEqual(
            clean_text("<p>Luxury <b>house</b> for sale</p>"),
            "Luxury house for sale",
        )

    def test_entities_are_decoded(self):
        from core.normalize import clean_text

        self.assertEqual(clean_text("Beds &amp; Wardrobes"), "Beds & Wardrobes")
        self.assertEqual(clean_text("Owner&#39;s car"), "Owner's car")
        self.assertEqual(clean_text("A&nbsp;B"), "A B")

    def test_plain_less_than_is_not_eaten(self):
        """'under < 50k' is prose, not markup."""
        from core.normalize import clean_text

        self.assertEqual(clean_text("Price under < 50k only"), "Price under < 50k only")
        self.assertEqual(clean_text("2 < 3 rooms"), "2 < 3 rooms")

    def test_ordinary_text_is_untouched(self):
        from core.normalize import clean_text

        self.assertEqual(clean_text("Honda Civic 2019 Oriel"), "Honda Civic 2019 Oriel")


class TestPrice(unittest.TestCase):
    def test_plain_and_grouped_numbers(self):
        self.assertEqual(parse_price("4500000")[0], Decimal("4500000"))
        self.assertEqual(parse_price("Rs 4,500,000")[0], Decimal("4500000"))
        self.assertEqual(parse_price("PKR 85,00,000")[0], Decimal("8500000"))

    def test_south_asian_magnitudes(self):
        self.assertEqual(parse_price("4.5 Crore")[0], Decimal("45000000"))
        self.assertEqual(parse_price("85 Lakh")[0], Decimal("8500000"))
        self.assertEqual(parse_price("1.2 crore")[0], Decimal("12000000"))

    def test_currency_detection(self):
        self.assertEqual(parse_price("USD 30,000")[1], "USD")
        self.assertEqual(parse_price("Rs 500")[1], "PKR")

    def test_no_usable_price(self):
        self.assertIsNone(parse_price("Call for price")[0])
        self.assertIsNone(parse_price("")[0])
        self.assertIsNone(parse_price(None)[0])

    def test_original_text_preserved(self):
        self.assertEqual(parse_price("Rs 4,500,000")[2], "Rs 4,500,000")


class TestAdDate(unittest.TestCase):
    def test_relative_phrases(self):
        self.assertEqual(parse_ad_date("3 days ago", now=NOW), date(2026, 9, 16))
        self.assertEqual(parse_ad_date("an hour ago", now=NOW), date(2026, 9, 19))
        self.assertEqual(parse_ad_date("Updated 2 weeks ago", now=NOW), date(2026, 9, 5))

    def test_today_and_yesterday(self):
        self.assertEqual(parse_ad_date("Today", now=NOW), date(2026, 9, 19))
        self.assertEqual(parse_ad_date("Yesterday", now=NOW), date(2026, 9, 18))

    def test_epoch_milliseconds(self):
        epoch_ms = int(datetime(2026, 5, 1, tzinfo=timezone.utc).timestamp() * 1000)
        self.assertEqual(parse_ad_date(epoch_ms), date(2026, 5, 1))

    def test_iso_string(self):
        self.assertEqual(parse_ad_date("2026-03-14T08:30:00Z"), date(2026, 3, 14))

    def test_absolute_date_matching_today_is_kept(self):
        """A listing posted today must not be discarded as a fuzzy-match artifact."""
        today = datetime.now(timezone.utc).date()
        text = today.strftime("%b %d, %Y")
        self.assertEqual(parse_ad_date(text), today)
        self.assertEqual(parse_ad_date(f"Last Updated: {text}"), today)

    def test_date_without_year_assumes_most_recent(self):
        result = parse_ad_date("Mar 14", now=NOW)
        self.assertEqual(result, date(2026, 3, 14))

    def test_year_rolls_back_when_date_would_be_future(self):
        result = parse_ad_date("Dec 25", now=NOW)  # NOW is September 2026
        self.assertEqual(result, date(2025, 12, 25))

    def test_text_without_a_date_returns_none(self):
        self.assertIsNone(parse_ad_date("Managed by PakWheels"))
        self.assertIsNone(parse_ad_date("Financing available"))

    def test_unparseable(self):
        self.assertIsNone(parse_ad_date(""))
        self.assertIsNone(parse_ad_date(None))


class TestPhone(unittest.TestCase):
    def test_normalizes_local_formats(self):
        for raw in ("03001234567", "0300-1234567", "+92 300 1234567", "00923001234567"):
            with self.subTest(raw=raw):
                self.assertEqual(normalize_phone(raw), "+923001234567")

    def test_rejects_junk(self):
        self.assertIsNone(normalize_phone("0000000000"))
        self.assertIsNone(normalize_phone("123"))
        self.assertIsNone(normalize_phone(""))

    def test_extracts_from_ad_text(self):
        text = "Serious buyers only. Contact 0321-9876543 after 6pm."
        self.assertEqual(extract_phone(text), "+923219876543")

    def test_returns_none_when_absent(self):
        self.assertIsNone(extract_phone("No numbers in this description at all."))

    def test_does_not_invent_numbers_from_spec_digits(self):
        """Regression: 'Mileage 41500 / 2019 import' was read as +92415002019."""
        self.assertIsNone(
            extract_phone("Mileage 41500\n2019 import/registered\n1.0l petrol")
        )
        self.assertIsNone(extract_phone("Price 2895000 Year 2021 Mileage 45000"))
        self.assertIsNone(extract_phone("Engine 1500cc, 4 seats, 3 owners"))

    def test_handles_arabic_indic_digits(self):
        self.assertEqual(extract_phone("رابطہ ٠٣٠٠١٢٣٤٥٦٧"), "+923001234567")

    def test_still_finds_real_numbers_around_noise(self):
        self.assertEqual(
            extract_phone("Mileage 45000\nCall 0300 1234567"), "+923001234567"
        )
        self.assertEqual(extract_phone("WhatsApp +92 321 9876543"), "+923219876543")
        self.assertEqual(extract_phone("Landline 042-35678901"), "+924235678901")


class TestListing(unittest.TestCase):
    def test_fingerprint_prefers_source_id(self):
        listing = make_listing()
        self.assertEqual(listing.fingerprint, "olx:id:1234567")

    def test_fingerprint_falls_back_to_url(self):
        listing = make_listing(source_listing_id=None)
        self.assertTrue(listing.fingerprint.startswith("olx:url:"))

    def test_canonical_url_strips_tracking(self):
        dirty = "https://www.olx.com.pk/item/abc-iid-1?utm_source=fb&fbclid=xyz"
        self.assertEqual(canonical_url(dirty), "https://olx.com.pk/item/abc-iid-1")

    def test_validation_accepts_good_record(self):
        self.assertEqual(make_listing().validate(), [])

    def test_validation_catches_problems(self):
        future = datetime.now(timezone.utc).date() + timedelta(days=5)
        self.assertTrue(make_listing(ad_date=future).validate())
        self.assertTrue(make_listing(title="").validate())
        self.assertTrue(make_listing(url="not-a-url").validate())
        self.assertTrue(make_listing(price=Decimal("-1")).validate())
        self.assertTrue(make_listing(source="craigslist").validate())


class TestPlanSelections(unittest.TestCase):
    """Per-source category selection, which the web UI sends."""

    def setUp(self):
        from core.config import load_config
        from core.pipeline import Pipeline

        self.config = load_config(PROJECT_ROOT)
        self.pipeline = Pipeline(self.config, db=None, http=None)

    def units(self, **kwargs):
        return [
            (c.name, s.name, cat.key)
            for c, s, cat in self.pipeline.plan(city_names=["Lahore"], **kwargs)
        ]

    def test_bikes_key_exists_on_two_sources(self):
        """The collision this whole feature exists to handle."""
        owners = {
            s.name
            for s in self.config.sources.values()
            for c in s.categories
            if c.key == "bikes"
        }
        self.assertEqual(owners, {"olx", "pakwheels"})

    def test_flat_category_list_cannot_separate_them(self):
        units = self.units(category_keys=["bikes"])
        self.assertEqual({s for _, s, _ in units}, {"olx", "pakwheels"})

    def test_selections_picks_one_source_only(self):
        units = self.units(selections={"olx": ["bikes"]})
        self.assertEqual(units, [("Lahore", "olx", "bikes")])

    def test_selections_across_sources(self):
        units = self.units(
            selections={"olx": ["cars"], "pakwheels": ["used-cars"]}
        )
        self.assertEqual(
            sorted(units),
            [("Lahore", "olx", "cars"), ("Lahore", "pakwheels", "used-cars")],
        )

    def test_empty_list_means_every_category_of_that_source(self):
        units = self.units(selections={"zameen": []})
        zameen = self.config.sources["zameen"]
        self.assertEqual(len(units), len(zameen.categories))
        self.assertTrue(all(s == "zameen" for _, s, _ in units))

    def test_unknown_source_is_rejected(self):
        with self.assertRaises(ValueError):
            self.units(selections={"craigslist": []})


class TestOlxCategoryTree(unittest.TestCase):
    """OLX's own menu tree - the source of the section grouping."""

    STATE = {
        "categories": {
            "data": [
                {
                    "name": "Mobiles",
                    "slug": "mobiles",
                    "externalID": "1411",
                    "level": 0,
                    "displayPriority": 100,
                    "children": [
                        {
                            "name": "Mobile Phones",
                            "slug": "mobile-phones",
                            "externalID": "1453",
                            "level": 1,
                            "displayPriority": 90,
                            "children": [
                                {
                                    "name": "Chargers",
                                    "slug": "mobile-chargers",
                                    "externalID": "1462",
                                    "level": 2,
                                }
                            ],
                        },
                        {
                            "name": "Tablets",
                            "slug": "tablets",
                            "externalID": "1455",
                            "level": 1,
                            "displayPriority": 80,
                        },
                    ],
                },
                {
                    "name": "Vehicles",
                    "slug": "vehicles",
                    "externalID": "5",
                    "level": 0,
                    "children": [
                        {"name": "Cars", "slug": "cars", "externalID": "84", "level": 1}
                    ],
                },
            ]
        }
    }

    def setUp(self):
        from core.categories import parse_category_tree

        self.categories = parse_category_tree(self.STATE)
        self.by_key = {c["key"]: c for c in self.categories}

    def test_every_node_is_captured(self):
        self.assertEqual(len(self.categories), 6)

    def test_children_inherit_their_top_level_section(self):
        self.assertEqual(self.by_key["mobile-phones"]["group"], "Mobiles")
        self.assertEqual(self.by_key["mobile-chargers"]["group"], "Mobiles")
        self.assertEqual(self.by_key["cars"]["group"], "Vehicles")

    def test_depth_is_recorded(self):
        self.assertEqual(self.by_key["mobiles"]["level"], 0)
        self.assertEqual(self.by_key["mobile-phones"]["level"], 1)
        self.assertEqual(self.by_key["mobile-chargers"]["level"], 2)

    def test_uses_olx_labels_not_slug_guesses(self):
        self.assertEqual(self.by_key["mobile-phones"]["label"], "Mobile Phones")
        self.assertEqual(self.by_key["mobile-chargers"]["label"], "Chargers")

    def test_path_is_built_from_slug_and_id(self):
        self.assertEqual(self.by_key["cars"]["path"], "cars_c84")

    def test_missing_tree_returns_nothing_so_the_sitemap_can_take_over(self):
        from core.categories import parse_category_tree

        self.assertEqual(parse_category_tree(None), [])
        self.assertEqual(parse_category_tree({}), [])
        self.assertEqual(parse_category_tree({"categories": {"data": []}}), [])

    def test_repeated_slug_across_sections_stays_unique(self):
        from core.categories import parse_category_tree

        state = {
            "categories": {
                "data": [
                    {
                        "name": "Property for Sale",
                        "slug": "property-for-sale",
                        "externalID": "2",
                        "children": [
                            {"name": "Houses", "slug": "houses", "externalID": "1719"}
                        ],
                    },
                    {
                        "name": "Property for Rent",
                        "slug": "property-for-rent",
                        "externalID": "3",
                        "children": [
                            {"name": "Houses", "slug": "houses", "externalID": "1721"}
                        ],
                    },
                ]
            }
        }
        keys = [c["key"] for c in parse_category_tree(state)]
        self.assertEqual(len(keys), len(set(keys)))
        self.assertIn("houses-c1719", keys)
        self.assertIn("houses-c1721", keys)


class TestOlxCategorySync(unittest.TestCase):
    SITEMAP = """
    <urlset>
      <url><loc>https://www.olx.com.pk/cars_c84</loc></url>
      <url><loc>https://www.olx.com.pk/audi-cars_c84</loc></url>
      <url><loc>https://www.olx.com.pk/bmw-cars_c84</loc></url>
      <url><loc>https://www.olx.com.pk/tv-video-audio_c729</loc></url>
      <url><loc>https://www.olx.com.pk/houses_c1719</loc></url>
      <url><loc>https://www.olx.com.pk/houses_c1721</loc></url>
      <url><loc>https://www.olx.com.pk/it-networking_c56</loc></url>
    </urlset>
    """

    def setUp(self):
        from core.categories import parse_categories

        self.categories = parse_categories(self.SITEMAP)
        self.by_key = {c["key"]: c for c in self.categories}

    def test_collapses_brand_variants_to_one_category(self):
        """audi-cars / bmw-cars / cars are all c84 - one category, not three."""
        cars = [c for c in self.categories if c["category_id"] == "84"]
        self.assertEqual(len(cars), 1)
        self.assertEqual(cars[0]["key"], "cars")
        self.assertEqual(cars[0]["path"], "cars_c84")

    def test_repeated_slug_across_ids_gets_disambiguated(self):
        """`houses` is both c1719 and c1721 - distinct categories."""
        self.assertIn("houses-c1719", self.by_key)
        self.assertIn("houses-c1721", self.by_key)
        self.assertEqual(self.by_key["houses-c1719"]["path"], "houses_c1719")
        self.assertEqual(self.by_key["houses-c1721"]["path"], "houses_c1721")

    def test_all_keys_are_unique(self):
        keys = [c["key"] for c in self.categories]
        self.assertEqual(len(keys), len(set(keys)))

    def test_labels_are_readable(self):
        from core.categories import humanize

        self.assertEqual(humanize("electronics-home-appliances"), "Electronics Home Appliances")
        self.assertEqual(humanize("tv-video-audio"), "TV Video Audio")
        self.assertEqual(humanize("it-networking"), "IT Networking")


class TestDateFilter(unittest.TestCase):
    """The pipeline keeps only ads posted inside the requested window."""

    def setUp(self):
        from core.pipeline import Pipeline, RunStats

        self._tmp = tempfile.TemporaryDirectory()
        self.db = Database(Path(self._tmp.name) / "t.db")
        self.pipeline = Pipeline.__new__(Pipeline)   # no config/http needed here
        self.pipeline.db = self.db
        self.stats = RunStats()

    def tearDown(self):
        self.db.close()
        self._tmp.cleanup()

    def collect(self, window, dates):
        """Run _collect_one against a stub collector yielding these ad dates."""
        class StubCollector:
            source_name = "olx"
            last_coverage = None
            last_outcome = None

            def collect(self, city, category, limit=None, date_window=(None, None),
                        stop_event=None):
                # date_window is accepted because a real collector uses it to
                # decide how deep to page; this stub yields a fixed set, so the
                # pipeline's own filtering is what is under test here.
                self.last_outcome = CityOutcome("Lahore", "olx", "Cars", status="completed")
                for index, ad_date in enumerate(dates):
                    yield make_listing(
                        source_listing_id=str(1000 + index),
                        url=f"https://www.olx.com.pk/item/x-iid-{1000 + index}",
                        ad_date=ad_date,
                    )

        city = type("C", (), {"name": "Lahore"})()
        category = type("K", (), {"key": "cars"})()
        self.pipeline._collect_one(
            StubCollector(), city, category, self.stats, None, None, date_window=window
        )
        return self.stats

    def stored_dates(self):
        rows = self.db.conn.execute("SELECT ad_date FROM listings").fetchall()
        return sorted((r["ad_date"] or "") for r in rows)

    def test_single_day_keeps_only_that_day(self):
        target = date(2026, 9, 20)
        self.collect(
            (target, target, True),
            [date(2026, 9, 19), target, target, date(2026, 9, 21)],
        )
        self.assertEqual(self.stored_dates(), ["2026-09-20", "2026-09-20"])
        self.assertEqual(self.stats.filtered, 2)

    def test_range_is_inclusive_at_both_ends(self):
        self.collect(
            (date(2026, 9, 18), date(2026, 9, 20), True),
            [date(2026, 9, 17), date(2026, 9, 18), date(2026, 9, 20), date(2026, 9, 21)],
        )
        self.assertEqual(self.stored_dates(), ["2026-09-18", "2026-09-20"])

    def test_open_ended_from(self):
        """No date_to means 'everything since'. Dates stay in the past because
        Listing.validate() rejects future ad dates."""
        recent = datetime.now(timezone.utc).date() - timedelta(days=1)
        cutoff = recent - timedelta(days=2)
        self.collect(
            (cutoff, None, True),
            [cutoff - timedelta(days=1), cutoff, recent],
        )
        self.assertEqual(
            self.stored_dates(), sorted([cutoff.isoformat(), recent.isoformat()])
        )

    def test_undated_listings_are_kept_by_default(self):
        """A parser gap must not silently drop real ads."""
        self.collect((date(2026, 9, 20), date(2026, 9, 20), True),
                     [None, date(2026, 9, 20)])
        self.assertEqual(len(self.stored_dates()), 2)
        self.assertEqual(self.stats.undated, 0)

    def test_undated_can_be_dropped_explicitly(self):
        self.collect((date(2026, 9, 20), date(2026, 9, 20), False),
                     [None, date(2026, 9, 20)])
        self.assertEqual(self.stored_dates(), ["2026-09-20"])
        self.assertEqual(self.stats.undated, 1)

    def test_no_window_keeps_everything(self):
        self.collect((None, None, True),
                     [date(2020, 1, 1), date(2026, 9, 20), None])
        self.assertEqual(len(self.stored_dates()), 3)
        self.assertEqual(self.stats.filtered, 0)


class TestDateParsing(unittest.TestCase):
    def test_accepts_iso_strings_and_dates(self):
        from core.pipeline import _as_date

        self.assertEqual(_as_date("2026-09-20"), date(2026, 9, 20))
        self.assertEqual(_as_date(date(2026, 9, 20)), date(2026, 9, 20))
        self.assertEqual(_as_date("2026-09-20T13:00:00"), date(2026, 9, 20))
        self.assertIsNone(_as_date(None))
        self.assertIsNone(_as_date(""))

    def test_rejects_nonsense(self):
        from core.pipeline import _as_date

        with self.assertRaises(ValueError):
            _as_date("20-09-2026")
        with self.assertRaises(ValueError):
            _as_date("yesterday")


class TestZameenPublishedPhone(unittest.TestCase):
    """Zameen renders a contact number on the public page; these are its shapes."""

    @staticmethod
    def extract(data: dict):
        from scrapers.zameen.collector import ZameenCollector

        return ZameenCollector._published_phone({"property": {"data": data}})

    def test_reads_primary_phone_number(self):
        self.assertEqual(
            self.extract({"primaryPhoneNumber": "+923373331572"}), "+923373331572"
        )

    def test_reads_nested_mobile_numbers_list(self):
        self.assertEqual(
            self.extract(
                {
                    "phoneNumber": {
                        "mobileNumbers": ["+923373331572"],
                        "phoneNumbers": ["+923373331572"],
                        "whatsapp": "923373331572",
                    }
                }
            ),
            "+923373331572",
        )

    def test_respects_login_gate(self):
        """When Zameen says a login is required, the number is not taken."""
        self.assertIsNone(
            self.extract(
                {
                    "requiresLoginForContact": True,
                    "primaryPhoneNumber": "+923373331572",
                }
            )
        )

    def test_returns_none_when_no_number_published(self):
        self.assertIsNone(self.extract({"contactName": "Yawar Abbas"}))
        self.assertIsNone(self.extract({"phoneNumber": {"proxyMobile": None}}))


class TestDedup(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)
        self.db = Database(self.tmp / "test.db")

    def tearDown(self):
        self.db.close()
        self._tmp.cleanup()

    def count(self) -> int:
        return self.db.conn.execute("SELECT COUNT(*) AS n FROM listings").fetchone()["n"]

    def test_same_listing_twice_stays_one_row(self):
        self.assertEqual(self.db.upsert(make_listing()), "inserted")
        self.assertEqual(self.db.upsert(make_listing()), "updated")
        self.assertEqual(self.count(), 1)

    def test_url_seen_before_id_is_not_duplicated(self):
        """An ad first seen without a visible id, then later with one."""
        self.db.upsert(make_listing(source_listing_id=None))
        result = self.db.upsert(make_listing(source_listing_id="1234567"))
        self.assertEqual(result, "updated")
        self.assertEqual(self.count(), 1)

        row = self.db.conn.execute("SELECT * FROM listings").fetchone()
        self.assertEqual(row["source_listing_id"], "1234567")
        self.assertEqual(row["fingerprint"], "olx:id:1234567")

    def test_different_sources_same_id_are_separate(self):
        self.db.upsert(make_listing(source="olx"))
        self.db.upsert(
            make_listing(source="zameen", url="https://www.zameen.com/Property/x-1234567.html")
        )
        self.assertEqual(self.count(), 2)

    def test_first_seen_preserved_last_seen_advances(self):
        self.db.upsert(make_listing())
        original = self.db.conn.execute("SELECT * FROM listings").fetchone()
        self.db.upsert(make_listing(title="Honda Civic 2019 - price reduced"))
        updated = self.db.conn.execute("SELECT * FROM listings").fetchone()

        self.assertEqual(original["first_seen_at"], updated["first_seen_at"])
        self.assertGreaterEqual(updated["last_seen_at"], original["last_seen_at"])
        self.assertEqual(updated["title"], "Honda Civic 2019 - price reduced")

    def test_richer_description_is_not_overwritten_by_thinner_one(self):
        self.db.upsert(make_listing(description="A" * 200))
        self.db.upsert(make_listing(description="short"))
        row = self.db.conn.execute("SELECT * FROM listings").fetchone()
        self.assertEqual(len(row["description"]), 200)

    def test_existing_phone_survives_a_scrape_without_one(self):
        self.db.upsert(make_listing(phone="+923001234567"))
        self.db.upsert(make_listing(phone=None))
        row = self.db.conn.execute("SELECT * FROM listings").fetchone()
        self.assertEqual(row["phone"], "+923001234567")


class TestExcelExport(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)
        self.db = Database(self.tmp / "test.db")
        self.db.upsert(make_listing(city="Lahore"))
        self.db.upsert(
            make_listing(
                city="Karachi",
                source="zameen",
                source_listing_id="99",
                url="https://www.zameen.com/Property/karachi-99.html",
            )
        )
        self.db.conn.commit()

    def tearDown(self):
        self.db.close()
        self._tmp.cleanup()

    def test_creates_workbook_with_city_sheets(self):
        from openpyxl import load_workbook

        path = export_to_excel(self.db, self.tmp / "out")
        self.assertTrue(path.exists())

        workbook = load_workbook(path)
        self.assertIn("Summary", workbook.sheetnames)
        self.assertIn("All Listings", workbook.sheetnames)
        self.assertIn("Lahore", workbook.sheetnames)
        self.assertIn("Karachi", workbook.sheetnames)

        sheet = workbook["All Listings"]
        self.assertEqual(sheet.max_row, 3)  # header + 2 listings

    def test_city_filter_limits_the_whole_workbook(self):
        """Regression: ?city=Lahore returned every city's rows and sheets."""
        from openpyxl import load_workbook

        path = export_to_excel(self.db, self.tmp / "out_city", city="Lahore")
        workbook = load_workbook(path)

        self.assertIn("Lahore", workbook.sheetnames)
        self.assertNotIn("Karachi", workbook.sheetnames)

        sheet = workbook["All Listings"]
        self.assertEqual(sheet.max_row, 2)  # header + the single Lahore listing

        cities = {
            row[3] for row in sheet.iter_rows(min_row=2, values_only=True)
        }
        self.assertEqual(cities, {"Lahore"})

    def test_seen_since_scopes_the_export_to_one_run(self):
        """Export what a run touched - new AND refreshed - not the whole table."""
        from openpyxl import load_workbook

        # An older listing that this run will not touch.
        self.db.upsert(
            make_listing(
                source_listing_id="old",
                url="https://www.olx.com.pk/item/old-iid-1",
                title="Collected last week",
            )
        )
        self.db.conn.execute(
            "UPDATE listings SET first_seen_at = ?, last_seen_at = ? "
            "WHERE source_listing_id = 'old'",
            ("2026-01-01T00:00:00+00:00", "2026-01-01T00:00:00+00:00"),
        )

        run_start = datetime.now(timezone.utc).isoformat()

        # One brand new, and one the run saw again.
        self.db.upsert(
            make_listing(
                source_listing_id="fresh",
                url="https://www.olx.com.pk/item/fresh-iid-2",
                title="Found in this run",
            )
        )
        self.db.upsert(make_listing(title="Seen again in this run"))
        self.db.conn.commit()

        path = export_to_excel(self.db, self.tmp / "out_run", seen_since=run_start)
        sheet = load_workbook(path)["All Listings"]
        title_at = [c.value for c in sheet[1]].index("Title")
        titles = {row[title_at] for row in sheet.iter_rows(min_row=2, values_only=True)}

        self.assertIn("Found in this run", titles)
        self.assertIn("Seen again in this run", titles)
        self.assertNotIn("Collected last week", titles)

    def test_description_has_a_preview_and_a_full_column(self):
        """Scannable preview next to the title, complete text further along."""
        from openpyxl import load_workbook

        long_text = (
            "Beautiful 3 bedroom house in a quiet street. " * 6
        ).strip()
        self.db.upsert(
            make_listing(
                source_listing_id="long",
                url="https://www.olx.com.pk/item/long-iid-9",
                description=long_text,
            )
        )
        self.db.conn.commit()

        path = export_to_excel(self.db, self.tmp / "out_desc")
        sheet = load_workbook(path)["All Listings"]
        headers = [c.value for c in sheet[1]]

        self.assertIn("Description", headers)
        self.assertIn("Full Description", headers)

        preview_at = headers.index("Description")
        full_at = headers.index("Full Description")

        row = next(
            r
            for r in sheet.iter_rows(min_row=2, values_only=True)
            if r[full_at] == long_text
        )

        self.assertTrue(row[preview_at].endswith("..."))
        self.assertLess(len(row[preview_at]), len(long_text))
        self.assertTrue(long_text.startswith(row[preview_at][:40]))

    def test_short_description_preview_is_not_truncated(self):
        from openpyxl import load_workbook

        path = export_to_excel(self.db, self.tmp / "out_short")
        sheet = load_workbook(path)["All Listings"]
        headers = [c.value for c in sheet[1]]
        preview_at = headers.index("Description")

        previews = [
            r[preview_at] for r in sheet.iter_rows(min_row=2, values_only=True) if r[preview_at]
        ]
        self.assertTrue(previews)
        self.assertFalse(any(p.endswith("...") for p in previews))

    def test_formula_injection_is_defused(self):
        from openpyxl import load_workbook

        self.db.upsert(
            make_listing(
                source_listing_id="777",
                url="https://www.olx.com.pk/item/evil-iid-777",
                description="=cmd|'/c calc'!A1",
            )
        )
        self.db.conn.commit()

        path = export_to_excel(self.db, self.tmp / "out2")
        sheet = load_workbook(path)["All Listings"]
        headers = [c.value for c in sheet[1]]
        preview_at = headers.index("Description")
        values = [
            row[preview_at]
            for row in sheet.iter_rows(min_row=2, values_only=True)
        ]
        injected = [v for v in values if v and "cmd" in str(v)]
        self.assertTrue(injected)
        self.assertTrue(all(str(v).startswith("'=") for v in injected))


if __name__ == "__main__":
    unittest.main(verbosity=2)
